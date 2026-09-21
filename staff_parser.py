import re
from openpyxl import load_workbook

AVAILABILITY_HEADERS = [
    'Ma -17','Ma >17','Di -17','Di >17','Wo -17','Wo >17','Do -17','Do >17',
    'Vr -17','Vr >17','Za -17','Za >17','Zo -17','Zo >17'
]
FIXED_HEADERS = {'Naam','Week','Opmerking','Rijbewijs','Eigen vervoer','Actief', *AVAILABILITY_HEADERS}


def _text(value):
    if value is None:
        return ''
    return re.sub(r'\s+', ' ', str(value)).strip()


def _yes_no_unknown(value):
    text = _text(value).casefold()
    if text in {'ja','j','yes','y','1','true','x','jaj'} or text.startswith('ja'):
        return 'Ja'
    if text in {'nee','n','no','0','false'} or text.startswith('nee'):
        return 'Nee'
    return 'Onbekend'


def _availability(value):
    text = _text(value)
    aliases = {
        '✓':'beschikbaar','✔':'beschikbaar','ja':'beschikbaar','beschikbaar':'beschikbaar',
        '~':'overleg','soms':'overleg','in overleg':'overleg','overleg':'overleg',
        '✕':'niet','×':'niet','x':'niet','nee':'niet','niet beschikbaar':'niet','niet':'niet',
        '?':'onbekend','onbekend':'onbekend','':'onbekend',
    }
    return aliases.get(text.casefold(), aliases.get(text, 'onbekend'))


def parse_personnel_workbook(fileobj):
    wb = load_workbook(fileobj, data_only=True, read_only=True)
    if 'Personeel' not in wb.sheetnames:
        raise ValueError("Tabblad 'Personeel' ontbreekt.")
    ws = wb['Personeel']

    header_row = None
    headers = []
    for r in range(1, min(ws.max_row, 20) + 1):
        row = [_text(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if 'Naam' in row and 'Ma -17' in row and 'Zo >17' in row:
            header_row = r
            headers = row
            break
    if not header_row:
        raise ValueError("Kopregel met Naam en weekbeschikbaarheid niet gevonden.")

    idx = {h: i for i, h in enumerate(headers) if h}
    missing = [h for h in ['Naam', *AVAILABILITY_HEADERS] if h not in idx]
    if missing:
        raise ValueError('Ontbrekende kolommen: ' + ', '.join(missing))

    skill_headers = [h for h in headers if h and h not in FIXED_HEADERS]
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        vals = [_text(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        name = vals[idx['Naam']] if idx.get('Naam') is not None and idx['Naam'] < len(vals) else ''
        if not name:
            continue
        week = vals[idx['Week']].upper() if 'Week' in idx and idx['Week'] < len(vals) else ''
        if week not in {'EVEN','ONEVEN'}:
            week = ''
        availability = {h: _availability(vals[idx[h]] if idx[h] < len(vals) else '') for h in AVAILABILITY_HEADERS}
        skills = {h: _yes_no_unknown(vals[idx[h]] if idx[h] < len(vals) else '') for h in skill_headers}
        rows.append({
            'name': name,
            'week_mode': week,
            'notes': vals[idx['Opmerking']] if 'Opmerking' in idx else '',
            'driving_license': vals[idx['Rijbewijs']] if 'Rijbewijs' in idx else '',
            'own_transport': _yes_no_unknown(vals[idx['Eigen vervoer']] if 'Eigen vervoer' in idx else ''),
            'active': _yes_no_unknown(vals[idx['Actief']] if 'Actief' in idx else 'Ja') != 'Nee',
            'availability': availability,
            'skills': skills,
        })
    if not rows:
        raise ValueError('Geen medewerkers gevonden in het bestand.')
    return rows, skill_headers
